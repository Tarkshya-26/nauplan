import sys, copy, subprocess, os, glob
from pptx import Presentation
from PIL import Image
src, prefix = sys.argv[1], sys.argv[2]
n = len(Presentation(src).slides)
os.makedirs('qa', exist_ok=True)
for k in range(n):
    p = Presentation(src)
    lst = p.slides._sldIdLst
    ids = list(lst)
    for i, s in enumerate(ids):
        if i != k: lst.remove(s)
    out = f'qa/{prefix}-{k+1}.pptx'; p.save(out)
    subprocess.run(['qlmanage','-t','-s','1600','-o','qa',out],capture_output=True)
    os.remove(out)
    os.replace(out+'.png', f'qa/{prefix}-{k+1}.png')
ims=[Image.open(f'qa/{prefix}-{k+1}.png') for k in range(n)]
w,h=800,450
g=Image.new('RGB',(w*2,h*((n+1)//2)),'white')
for i,im in enumerate(ims): g.paste(im.resize((w,h)),((i%2)*w,(i//2)*h))
g.save(f'qa/{prefix}-grid.png'); print('ok',n)
