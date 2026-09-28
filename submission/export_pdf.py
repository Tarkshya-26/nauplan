"""Render each slide with macOS Quick Look and assemble a 6-page PDF (image based)."""
import os
import subprocess

from PIL import Image
from pptx import Presentation

SRC = "SIH2026_NauPlan_Vector66.pptx"
n = len(Presentation(SRC).slides)
os.makedirs("pdfq", exist_ok=True)
pages = []
for k in range(n):
    p = Presentation(SRC)
    lst = p.slides._sldIdLst
    for i, s in enumerate(list(lst)):
        if i != k:
            lst.remove(s)
    out = f"pdfq/s{k + 1}.pptx"
    p.save(out)
    subprocess.run(["qlmanage", "-t", "-s", "3200", "-o", "pdfq", out], capture_output=True)
    pages.append(Image.open(out + ".png").convert("RGB"))
    os.remove(out)
pages[0].save("SIH2026_NauPlan_Vector66.pdf", save_all=True, append_images=pages[1:], resolution=240)
print(f"{n} slides -> SIH2026_NauPlan_Vector66.pdf")
